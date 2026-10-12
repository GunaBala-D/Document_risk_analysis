"""Text extraction with OCR fallback.

PDF   -> PyMuPDF native text layer, page by page. Pages whose text layer is missing or
         too thin (scanned pages) are rasterised and OCR'd with Tesseract.
DOCX  -> python-docx (paragraphs + tables).
Image -> Tesseract OCR.

OCR failures never crash the pipeline: they produce a result with
``status="failed"`` and a clear error so the API can report it.
"""

from __future__ import annotations

import io
import time
from dataclasses import dataclass, field

import pymupdf as fitz

from src.logging_utils import get_logger

log = get_logger(__name__)

MIN_CHARS_PER_PAGE = 50  # below this a PDF page is treated as scanned
OCR_DPI = 300


class OCRUnavailableError(RuntimeError):
    pass


@dataclass
class ExtractionResult:
    text: str
    status: str  # success | partial | failed
    method: str  # native | ocr | mixed | docx
    page_count: int
    ocr_pages: list[int] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    error: str | None = None
    duration_ms: float = 0.0

    @property
    def char_count(self) -> int:
        return len(self.text)


def _configure_tesseract(tesseract_cmd: str | None) -> None:
    import pytesseract

    if tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = tesseract_cmd


def ocr_image_bytes(image_bytes: bytes, tesseract_cmd: str | None = None) -> str:
    """OCR a PNG/JPEG byte string. Raises OCRUnavailableError if Tesseract is missing."""
    import pytesseract
    from PIL import Image

    _configure_tesseract(tesseract_cmd)
    try:
        with Image.open(io.BytesIO(image_bytes)) as img:
            gray = img.convert("L")
            # --psm 6: assume a uniform block of text; good for contract pages.
            return pytesseract.image_to_string(gray, lang="eng", config="--psm 6")
    except pytesseract.TesseractNotFoundError as exc:
        raise OCRUnavailableError("Tesseract binary not found; set TESSERACT_CMD.") from exc


def _extract_pdf(data: bytes, tesseract_cmd: str | None) -> ExtractionResult:
    pages_text: list[str] = []
    ocr_pages: list[int] = []
    warnings: list[str] = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        page_count = doc.page_count
        for i, page in enumerate(doc):
            native = page.get_text("text") or ""
            if len(native.strip()) >= MIN_CHARS_PER_PAGE:
                pages_text.append(native)
                continue
            try:
                pix = page.get_pixmap(dpi=OCR_DPI, colorspace=fitz.csGRAY)
                pages_text.append(ocr_image_bytes(pix.tobytes("png"), tesseract_cmd))
                ocr_pages.append(i + 1)
            except OCRUnavailableError:
                raise
            except Exception as exc:  # one bad page must not lose the whole document
                warnings.append(f"OCR failed on page {i + 1}: {exc}")
                pages_text.append(native)
    if not ocr_pages:
        method = "native"
    elif len(ocr_pages) == page_count:
        method = "ocr"
    else:
        method = "mixed"
    return ExtractionResult(
        text="\n".join(pages_text), status="partial" if warnings else "success",
        method=method, page_count=page_count, ocr_pages=ocr_pages, warnings=warnings,
    )


def _extract_docx(data: bytes) -> ExtractionResult:
    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return ExtractionResult(text="\n".join(parts), status="success", method="docx", page_count=1)


def _extract_image(data: bytes, tesseract_cmd: str | None) -> ExtractionResult:
    text = ocr_image_bytes(data, tesseract_cmd)
    return ExtractionResult(text=text, status="success", method="ocr", page_count=1, ocr_pages=[1])


def extract_text(data: bytes, file_format: str, tesseract_cmd: str | None = None) -> ExtractionResult:
    start = time.perf_counter()
    try:
        if file_format == "pdf":
            result = _extract_pdf(data, tesseract_cmd)
        elif file_format == "docx":
            result = _extract_docx(data)
        elif file_format in {"png", "jpeg"}:
            result = _extract_image(data, tesseract_cmd)
        else:
            raise ValueError(f"Unsupported format: {file_format}")
    except OCRUnavailableError as exc:
        result = ExtractionResult(text="", status="failed", method="ocr", page_count=0, error=str(exc))
    except Exception as exc:
        log.exception("text_extraction_failed", extra={"file_format": file_format})
        result = ExtractionResult(text="", status="failed", method="unknown", page_count=0,
                                  error=f"{type(exc).__name__}: {exc}")

    if result.status != "failed" and len(result.text.strip()) < MIN_CHARS_PER_PAGE:
        result.status = "failed"
        result.error = "No readable text could be extracted (blank or unreadable document)."
    result.duration_ms = round((time.perf_counter() - start) * 1000, 2)
    return result
